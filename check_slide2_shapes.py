import sys, os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\AI\dự thi.pptx')
slide2 = prs.slides[1]

for idx, s in enumerate(slide2.shapes):
    print(f"Shape {idx}: ID={s.shape_id}, Name='{s.name}', Type={s.shape_type}")
    if s.shape_type == 6: # Group
        for sub in s.shapes:
            print(f"   Subshape: ID={sub.shape_id}, Name='{sub.name}', Type={sub.shape_type}")
            if sub.has_text_frame:
                print(f"     Text: '{sub.text_frame.text.strip()}'")
