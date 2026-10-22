import sys, os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\AI\dự thi.pptx')

for idx in [1, 2]: # Slide 2 and Slide 3
    slide = prs.slides[idx]
    print(f"\n=== SLIDE {idx+1} ===")
    for s in slide.shapes:
        txt = s.text_frame.text.strip().replace('\n', ' ') if s.has_text_frame else 'N/A'
        print(f"Shape: id={s.shape_id}, name='{s.name}', L={s.left/914400:.2f}, T={s.top/914400:.2f}, W={s.width/914400:.2f}, H={s.height/914400:.2f}, Text='{txt}'")
