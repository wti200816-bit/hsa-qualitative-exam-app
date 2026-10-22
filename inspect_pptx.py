import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding='utf-8')

def inspect(path):
    print(f"\n{'='*30}\nFILE: {path}\n{'='*30}")
    if not os.path.exists(path):
        print("File does not exist!")
        return
    prs = Presentation(path)
    print(f"Total slides: {len(prs.slides)}")
    print(f"Dimensions: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} inches")
    for i, slide in enumerate(prs.slides):
        print(f"\n--- SLIDE {i+1} ---")
        for j, shape in enumerate(slide.shapes):
            pos_str = f"L={shape.left/914400:.2f} T={shape.top/914400:.2f} W={shape.width/914400:.2f} H={shape.height/914400:.2f}"
            st = shape.shape_type
            if shape.has_text_frame:
                txt = shape.text_frame.text.replace("\n", " \\n ")
                print(f"  Shape [{j}] ID={shape.shape_id} Name='{shape.name}' Pos=({pos_str})")
                print(f"     TEXT: {txt}")
            elif shape.shape_type == 13: # Picture
                print(f"  Shape [{j}] ID={shape.shape_id} Name='{shape.name}' [PICTURE] Pos=({pos_str})")
            else:
                print(f"  Shape [{j}] ID={shape.shape_id} Name='{shape.name}' Type={st} Pos=({pos_str})")

if __name__ == '__main__':
    inspect(r'E:\AI\dự thi.pptx')
    inspect(r'C:\Users\IASNNNN\Downloads\Mẫu Slide 2026.pptx')
