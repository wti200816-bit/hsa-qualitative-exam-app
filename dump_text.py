import sys, os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\AI\dự thi.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            print(f"--- Shape {s_idx}: Name='{shape.name}', ID={shape.shape_id} ---")
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                text = p.text.strip()
                if text:
                    print(f"  P[{p_idx}]: {text}")
