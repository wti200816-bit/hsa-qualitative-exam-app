import os
import sys
import shutil
import qrcode
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# ==========================================
# PALETTE MÀU SẮC CHUẨN BAN TỔ CHỨC
# ==========================================
COLOR_NAVY_DARK = RGBColor(0, 32, 96)       # Xanh đậm tiêu đề BTC (#002060)
COLOR_NAVY_BLUE = RGBColor(0, 70, 160)      # Xanh biển đậm (#0046A0)
COLOR_TEXT_MAIN = RGBColor(38, 38, 38)      # Màu chữ chính (#262626)
COLOR_TEXT_MUTED = RGBColor(90, 90, 90)     # Màu chữ phụ (#5A5A5A)
COLOR_TEXT_WHITE = RGBColor(255, 255, 255)  # Trắng
COLOR_GOLD = RGBColor(255, 204, 0)          # Vàng cam nổi bật
COLOR_ACCENT_RED = RGBColor(192, 0, 0)      # Đỏ đô logo RMIT/BGD
COLOR_BG_CARD = RGBColor(248, 250, 253)     # Nền card xám xanh rất nhạt
COLOR_BORDER_CARD = RGBColor(200, 215, 230) # Viền card

FONT_MAIN = 'Arial'

def set_font(run, name=FONT_MAIN, size_pt=15, bold=False, italic=False, color=COLOR_TEXT_MAIN):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color

def create_card(slide, left, top, width, height, bg_color=COLOR_BG_CARD, border_color=COLOR_BORDER_CARD, border_width_pt=1.5):
    """Tạo một hình khối chữ nhật bo góc làm card nền"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width_pt)
    else:
        shape.line.fill.background()
    return shape

def add_header_box(slide, left, top, width, height, title_text, bg_color=COLOR_NAVY_DARK, text_color=COLOR_TEXT_WHITE, font_size=15):
    """Tạo thanh tiêu đề cho card"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_MAIN
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = text_color
    return shape

def clear_shape_text(shape):
    """Xóa sạch text trong một shape"""
    if shape.has_text_frame:
        tf = shape.text_frame
        for p in list(tf.paragraphs):
            p.text = ""

def generate_qr_images():
    os.makedirs(r'E:\AI', exist_ok=True)
    qr1_path = r'E:\AI\qr_webapp.png'
    qr2_path = r'E:\AI\qr_github.png'
    if not os.path.exists(qr1_path):
        qr1 = qrcode.QRCode(version=1, box_size=10, border=2)
        qr1.add_data('https://transcendent-cuchufli-30f705.netlify.app')
        qr1.make(fit=True)
        img1 = qr1.make_image(fill_color='#002060', back_color='white')
        img1.save(qr1_path)
    if not os.path.exists(qr2_path):
        qr2 = qrcode.QRCode(version=1, box_size=10, border=2)
        qr2.add_data('https://wti200816-bit.github.io/hsa-qualitative-exam-app/')
        qr2.make(fit=True)
        img2 = qr2.make_image(fill_color='#002060', back_color='white')
        img2.save(qr2_path)
    return qr1_path, qr2_path

def build_presentation(input_pptx, output_pptx):
    print(f"Loading presentation from: {input_pptx}")
    prs = Presentation(input_pptx)
    print(f"Initial slide count: {len(prs.slides)}")

    qr_webapp, qr_github = generate_qr_images()
    app_screenshot = r'E:\AI\app_screenshot.png'

    # ====================================================
    # SLIDE 1: TRANG BÌA
    # ====================================================
    print("Formatting Slide 1 (Trang bìa)...")
    slide1 = prs.slides[0]
    for s in slide1.shapes:
        if s.name == 'TextBox 11':
            s.top = Inches(4.0)
            s.height = Inches(2.2)
            tf = s.text_frame
            tf.clear()
            tf.word_wrap = True
            p1 = tf.paragraphs[0]
            p1.text = "HSA Master AI: Giải pháp số hóa học liệu Ngữ Văn"
            p1.alignment = PP_ALIGN.CENTER
            p1.line_spacing = 1.15
            set_font(p1.runs[0], size_pt=30, bold=True, color=COLOR_NAVY_DARK)
            
            p2 = tf.add_paragraph()
            p2.text = "và Gia sư ảo hỗ trợ học sinh THPT chinh phục kỳ thi Đánh giá năng lực"
            p2.alignment = PP_ALIGN.CENTER
            p2.line_spacing = 1.15
            set_font(p2.runs[0], size_pt=26, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 12':
            tf = s.text_frame
            tf.clear()
            tf.word_wrap = True
            p1 = tf.paragraphs[0]
            p1.text = "Đơn vị công tác:"
            set_font(p1.runs[0], size_pt=18, bold=True, color=COLOR_TEXT_WHITE)
            p2 = tf.add_paragraph()
            p2.text = "Tổ Ngữ văn – Ngoại ngữ | Trường THPT Chương Mỹ A"
            set_font(p2.runs[0], size_pt=16, bold=False, color=COLOR_TEXT_WHITE)
        elif s.name == 'TextBox 13':
            tf = s.text_frame
            tf.clear()
            tf.word_wrap = True
            p1 = tf.paragraphs[0]
            p1.text = "Email:"
            set_font(p1.runs[0], size_pt=18, bold=True, color=COLOR_TEXT_WHITE)
            p2 = tf.add_paragraph()
            p2.text = "Tieulanthuytinh@gmail.com"
            set_font(p2.runs[0], size_pt=16, bold=False, color=COLOR_TEXT_WHITE)

    # ====================================================
    # SLIDE 2: NỘI DUNG BÀI TRÌNH BÀY (DÀN Ý + ẢNH MẪU WEB APP)
    # ====================================================
    print("Formatting Slide 2 (Nội dung bài trình bày)...")
    slide2 = prs.slides[1]
    
    # Ẩn/xóa text cũ
    for s in slide2.shapes:
        if s.name == 'TextBox 7':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "NỘI DUNG BÀI TRÌNH BÀY"
            set_font(p.runs[0], size_pt=38, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 9' or s.name == 'TextBox 10':
            clear_shape_text(s)
            
    # Chèn ảnh Web App vào khung bên trái
    if os.path.exists(app_screenshot):
        img_left = Inches(1.25)
        img_top = Inches(1.35)
        img_width = Inches(7.34)
        img_height = Inches(4.55)
        slide2.shapes.add_picture(app_screenshot, img_left, img_top, width=img_width, height=img_height)
        
        # Caption phía dưới ảnh
        caption_box = slide2.shapes.add_textbox(Inches(1.25), Inches(6.15), Inches(7.34), Inches(2.2))
        ctf = caption_box.text_frame
        ctf.word_wrap = True
        cp1 = ctf.paragraphs[0]
        cp1.text = "Giao diện Web App HSA Master AI"
        cp1.alignment = PP_ALIGN.CENTER
        set_font(cp1.runs[0], size_pt=16, bold=True, color=COLOR_NAVY_DARK)
        cp2 = ctf.add_paragraph()
        cp2.text = "• Giao diện chia đôi màn hình (Split-Screen) đọc hiểu văn bản dài\n• 50 đề chuẩn hóa ma trận HSA & 4.028 câu hỏi số hóa\n• Gia sư ảo AI chấm điểm & giải thích 1-1 tức thì 24/7"
        cp2.line_spacing = 1.25
        for r in cp2.runs:
            set_font(r, size_pt=13.5, color=COLOR_TEXT_MUTED)

    # Khung Card nền bên phải chứa 6 mục
    create_card(slide2, Inches(10.2), Inches(2.1), Inches(8.8), Inches(7.2))
    
    # 6 mục dàn ý bên phải
    agenda_items = [
        ("1. Bối cảnh và vấn đề:", "Nhu cầu ôn thi HSA & 3 rào cản tại THPT Chương Mỹ A."),
        ("2. Ý tưởng và giải pháp:", "Web App HSA Master AI ứng dụng mô hình \"AI Kép có trách nhiệm\"."),
        ("3. Quy trình thực hiện:", "Kết hợp Antigravity IDE, GitHub, MCP Tavily & Gemini AI."),
        ("4. Triển khai thực tế:", "Thử nghiệm tại lớp 12A2 và 12A7 (Tháng 10/2026)."),
        ("5. Kết quả và tác động:", "Cải thiện +24,7% điểm số, hỗ trợ phân tích chuyên đề yếu."),
        ("6. Thảo luận, tổng kết & Tài liệu:", "Tính bền vững (0đ chi phí) và khả năng nhân rộng.")
    ]
    
    # Clear các TextBox cũ của BTC ở bên phải
    for s in slide2.shapes:
        if s.shape_id in [12, 14, 16, 18, 20, 22]:
            clear_shape_text(s)

    # Tạo 6 card con xếp gọn gàng trong Card bên phải
    start_top = Inches(2.3)
    item_height = Inches(1.05)
    item_gap = Inches(0.12)
    
    for i, (title, desc) in enumerate(agenda_items):
        cur_top = start_top + i * (item_height + item_gap)
        sub_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.4), cur_top, Inches(8.4), item_height)
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sub_card.line.color.rgb = RGBColor(220, 230, 242)
        sub_card.line.width = Pt(1)
        
        badge = slide2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.55), cur_top + Inches(0.18), Inches(0.68), Inches(0.68))
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_NAVY_DARK if i < 4 else COLOR_ACCENT_RED
        badge.line.fill.background()
        btf = badge.text_frame
        btf.margin_top = Inches(0.12)
        bp = btf.paragraphs[0]
        bp.text = str(i + 1)
        bp.alignment = PP_ALIGN.CENTER
        set_font(bp.runs[0], size_pt=14, bold=True, color=COLOR_TEXT_WHITE)

        tb = slide2.shapes.add_textbox(Inches(11.35), cur_top + Inches(0.08), Inches(7.35), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.line_spacing = 1.15
        r1 = p1.add_run()
        r1.text = title.split(".", 1)[1].strip() + "\n"
        set_font(r1, size_pt=14.5, bold=True, color=COLOR_NAVY_DARK)
        
        r2 = p1.add_run()
        r2.text = desc
        set_font(r2, size_pt=13, bold=False, color=COLOR_TEXT_MAIN)

    # ====================================================
    # XÓA SLIDE 3 (Slide lời dặn quy chế BTC) để thành 9 slide chuẩn
    # ====================================================
    print("Removing BTC guideline slide (Slide 3)...")
    rId = prs.slides._sldIdLst[2].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[2]
    print(f"Slide count after removal: {len(prs.slides)}")

    # ====================================================
    # SLIDE 3 (mới): BỐI CẢNH VÀ VẤN ĐỀ
    # ====================================================
    print("Formatting Slide 3 (Bối cảnh và vấn đề)...")
    slide3 = prs.slides[2]
    for s in slide3.shapes:
        if s.name == 'TextBox 3':
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "BỐI CẢNH VÀ VẤN ĐỀ"
            set_font(p.runs[0], size_pt=40, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 4':
            clear_shape_text(s)

    # Card Trái: Vấn đề thực tế
    create_card(slide3, Inches(1.0), Inches(2.1), Inches(8.7), Inches(7.1))
    add_header_box(slide3, Inches(1.0), Inches(2.1), Inches(8.7), Inches(0.55), "1. THỰC TRẠNG & RÀO CẢN THỰC TẾ", bg_color=COLOR_NAVY_DARK)
    
    tb_left = slide3.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(8.3), Inches(6.2))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True
    
    bullets_l = [
        ("Kỳ thi ĐGNL ĐHQGHN (HSA) đặc biệt quan trọng: ", "Là phương thức xét tuyển đại học hàng đầu của học sinh khối 12 THPT Chương Mỹ A. Phần thi Ngôn ngữ – Văn học chiếm 1/3 tổng điểm (50 câu/60 phút) với nhiều ngữ liệu văn bản rất dài."),
        ("Không thuộc chương trình dạy chính khóa: ", "Thời lượng giảng dạy trên lớp phải tập trung cho chương trình chuẩn phổ thông; giáo viên và học sinh không có thời lượng ôn luyện chuyên biệt, bắt buộc học sinh phải tự học ngoài giờ."),
        ("Thiếu công cụ tự học định hướng: ", "Đề giấy hoặc Google Forms cuộn trang mỏi mắt với nhóm câu hỏi chùm; chỉ có đáp án trắc nghiệm A/B/C/D mà hoàn toàn không có lời giải thích sâu vì sao làm sai."),
        ("Chi phí ôn thi đắt đỏ: ", "Các khóa luyện thi HSA tại các trung tâm thường thu học phí rất cao, tạo áp lực kinh tế nặng nề cho học sinh và phụ huynh khu vực ngoại thành.")
    ]
    
    for i, (head, body) in enumerate(bullets_l):
        p = tf_l.paragraphs[0] if i == 0 else tf_l.add_paragraph()
        p.line_spacing = 1.2
        p.space_after = Pt(12)
        r0 = p.add_run()
        r0.text = "• "
        set_font(r0, size_pt=15, bold=True, color=COLOR_NAVY_BLUE)
        r1 = p.add_run()
        r1.text = head
        set_font(r1, size_pt=15, bold=True, color=COLOR_NAVY_DARK)
        r2 = p.add_run()
        r2.text = body
        set_font(r2, size_pt=14.5, bold=False, color=COLOR_TEXT_MAIN)

    # Card Phải: Số liệu khảo sát thực tế tại 12A2 & 12A7
    create_card(slide3, Inches(10.3), Inches(2.1), Inches(8.7), Inches(7.1))
    add_header_box(slide3, Inches(10.3), Inches(2.1), Inches(8.7), Inches(0.55), "2. SỐ LIỆU KHẢO SÁT THỰC TẾ (LỚP 12A2 & 12A7)", bg_color=COLOR_ACCENT_RED)
    
    tb_right = slide3.shapes.add_textbox(Inches(10.5), Inches(2.8), Inches(8.3), Inches(6.2))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True
    
    stats_r = [
        ("> 80% Học sinh có nguyện vọng thi HSA: ", "Qua khảo sát thực tế tại 2 lớp 12A2 & 12A7, hơn 80% học sinh xác định HSA là mục tiêu xét tuyển đại học trọng điểm nhưng hoàn toàn thiếu học liệu chuẩn."),
        ("78% Gặp khó khăn đọc dài & lặp lại lỗi sai: ", "Học sinh rất dễ nản lòng khi gặp các văn bản đọc hiểu chùm dài; thường xuyên lặp lại lỗi sai ở phần Tiếng Việt & Ngữ pháp do không có người giải thích tức thì ngoài giờ học."),
        ("Rủi ro khi tự dùng ChatGPT ngoài luồng: ", "Học sinh mất nhiều thời gian chụp/copy-paste đề; ChatGPT tự do dễ bịa đặt kiến thức văn học (ảo giác AI) hoặc làm lộ đáp án, gây tâm lý ỷ lại và học vẹt.")
    ]
    
    for i, (head, body) in enumerate(stats_r):
        p = tf_r.paragraphs[0] if i == 0 else tf_r.add_paragraph()
        p.line_spacing = 1.25
        p.space_after = Pt(14)
        r0 = p.add_run()
        r0.text = "• "
        set_font(r0, size_pt=15, bold=True, color=COLOR_ACCENT_RED)
        r1 = p.add_run()
        r1.text = head
        set_font(r1, size_pt=15.5, bold=True, color=COLOR_ACCENT_RED)
        r2 = p.add_run()
        r2.text = body
        set_font(r2, size_pt=14.5, bold=False, color=COLOR_TEXT_MAIN)

    # ====================================================
    # SLIDE 4 (mới): Ý TƯỞNG VÀ GIẢI PHÁP - PHẦN 1
    # ====================================================
    print("Formatting Slide 4 (Ý tưởng & Giải pháp - Phần 1)...")
    slide4 = prs.slides[3]
    for s in slide4.shapes:
        if s.name == 'TextBox 6':
            s.width = Inches(18.0)
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "Ý TƯỞNG VÀ GIẢI PHÁP: MÔ HÌNH \"AI KÉP CÓ TRÁCH NHIỆM\""
            set_font(p.runs[0], size_pt=34, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 2':
            clear_shape_text(s)

    # Banner trên
    create_card(slide4, Inches(1.0), Inches(2.0), Inches(18.0), Inches(1.4), bg_color=RGBColor(240, 246, 255), border_color=COLOR_NAVY_BLUE)
    tb_b = slide4.shapes.add_textbox(Inches(1.2), Inches(2.05), Inches(17.6), Inches(1.3))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    bp1 = tf_b.paragraphs[0]
    bp1.line_spacing = 1.2
    br1 = bp1.add_run()
    br1.text = "Giải pháp đột phá: "
    set_font(br1, size_pt=16, bold=True, color=COLOR_NAVY_DARK)
    br2 = bp1.add_run()
    br2.text = "Xây dựng Web App chuyên biệt \"HSA Master AI\" giải quyết toàn diện bài toán tự học Ngữ Văn."
    set_font(br2, size_pt=16, bold=False, color=COLOR_TEXT_MAIN)
    bp2 = tf_b.add_paragraph()
    bp2.line_spacing = 1.2
    br3 = bp2.add_run()
    br3.text = "Công cụ & Nền tảng: "
    set_font(br3, size_pt=15, bold=True, color=COLOR_ACCENT_RED)
    br4 = bp2.add_run()
    br4.text = "Google Antigravity IDE (Agentic AI), GitHub & GitHub Pages, MCP Tavily Search và Gemini API."
    set_font(br4, size_pt=15, bold=False, color=COLOR_TEXT_MAIN)

    # Cột Trái: Giá trị vượt trội
    create_card(slide4, Inches(1.0), Inches(3.6), Inches(8.8), Inches(5.6))
    add_header_box(slide4, Inches(1.0), Inches(3.6), Inches(8.8), Inches(0.55), "GIÁ TRỊ VƯỢT TRỘI SO VỚI TRUYỀN THỐNG", bg_color=COLOR_NAVY_DARK)
    
    tb_s4_l = slide4.shapes.add_textbox(Inches(1.2), Inches(4.3), Inches(8.4), Inches(4.7))
    tf_s4_l = tb_s4_l.text_frame
    tf_s4_l.word_wrap = True
    
    feats_l = [
        ("Kho đề trộn tự động chuẩn ma trận ĐHQGHN: ", "Hệ thống tự động bốc tách, phân loại câu hỏi theo 4 mức độ nhận thức (Nhận biết – Thông hiểu – Vận dụng – Vận dụng cao), tạo đề thi phong phú không trùng lặp."),
        ("Giao diện Chia đôi màn hình (Split-Screen): ", "Ngữ liệu bài đọc dài được cố định bên trái, nhóm câu hỏi chùm hiển thị bên phải. Học sinh dễ dàng đối chiếu trực tiếp, khắc phục hoàn toàn nhược điểm cuộn trang mỏi mắt của Google Forms."),
        ("Gia sư AI (Context-Aware AI Tutor): ", "Tự động nạp sẵn ngữ cảnh văn bản và câu hỏi sai để phân tích cặn kẽ 1-1 ngay sau khi nộp bài, giúp học sinh sửa lỗi sai lập tức mà không cần chờ giáo viên.")
    ]
    for i, (head, body) in enumerate(feats_l):
        p = tf_s4_l.paragraphs[0] if i == 0 else tf_s4_l.add_paragraph()
        p.line_spacing = 1.2
        p.space_after = Pt(10)
        r0 = p.add_run()
        r0.text = "• "
        set_font(r0, size_pt=15, bold=True, color=COLOR_NAVY_BLUE)
        r1 = p.add_run()
        r1.text = head
        set_font(r1, size_pt=15, bold=True, color=COLOR_NAVY_DARK)
        r2 = p.add_run()
        r2.text = body
        set_font(r2, size_pt=14.5, bold=False, color=COLOR_TEXT_MAIN)

    # Cột Phải: Đạo đức & AI có trách nhiệm
    create_card(slide4, Inches(10.2), Inches(3.6), Inches(8.8), Inches(5.6))
    add_header_box(slide4, Inches(10.2), Inches(3.6), Inches(8.8), Inches(0.55), "ĐẠO ĐỨC & TRÍ TUỆ NHÂN TẠO CÓ TRÁCH NHIỆM", bg_color=COLOR_ACCENT_RED)
    
    tb_s4_r = slide4.shapes.add_textbox(Inches(10.4), Inches(4.3), Inches(8.4), Inches(4.7))
    tf_s4_r = tb_s4_r.text_frame
    tf_s4_r.word_wrap = True
    
    feats_r = [
        ("Chống AI bịa đặt kiến thức (Anti-Hallucination): ", "Sử dụng 100% đề gốc do Tổ chuyên môn biên soạn. AI đóng vai trò tra cứu và giải bổ sung các câu khuyết dưới sự kiểm duyệt chặt chẽ của giáo viên trước khi phát hành."),
        ("Chống gian lận học thuật: ", "Tự động khóa hoàn toàn Gia sư AI trong suốt 60 phút học sinh làm bài thi thử bấm giờ; chỉ mở khóa khi học sinh luyện tập chuyên đề tự do hoặc sau khi đã nộp bài để xem phân tích."),
        ("Khuyến khích tư duy phản biện: ", "AI không chỉ đưa ra đáp án đúng mà tập trung giải thích bản chất câu sai, giúp học sinh phát triển năng lực tự học và tư duy ngôn ngữ sâu sắc.")
    ]
    for i, (head, body) in enumerate(feats_r):
        p = tf_s4_r.paragraphs[0] if i == 0 else tf_s4_r.add_paragraph()
        p.line_spacing = 1.2
        p.space_after = Pt(10)
        r0 = p.add_run()
        r0.text = "• "
        set_font(r0, size_pt=15, bold=True, color=COLOR_ACCENT_RED)
        r1 = p.add_run()
        r1.text = head
        set_font(r1, size_pt=15, bold=True, color=COLOR_ACCENT_RED)
        r2 = p.add_run()
        r2.text = body
        set_font(r2, size_pt=14.5, bold=False, color=COLOR_TEXT_MAIN)

    # ====================================================
    # SLIDE 5 (mới): Ý TƯỞNG VÀ GIẢI PHÁP - PHẦN 2 (CÁCH THỰC HIỆN)
    # ====================================================
    print("Formatting Slide 5 (Quy trình thực hiện 5 bước)...")
    slide5 = prs.slides[4]
    for s in slide5.shapes:
        if s.name == 'TextBox 6':
            s.width = Inches(18.0)
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "QUY TRÌNH THỰC HIỆN DỰ ÁN (5 BƯỚC CHUẨN HÓA)"
            set_font(p.runs[0], size_pt=36, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 2':
            clear_shape_text(s)

    steps = [
        ("BƯỚC 1", "Chuẩn bị ngân hàng học liệu chuẩn ma trận", "Xây dựng hệ thống câu hỏi theo đúng ma trận cấu trúc đề thi mẫu HSA của ĐHQG Hà Nội (50 câu/60 phút: Tiếng Việt, Đọc hiểu văn bản và Văn học).", COLOR_NAVY_DARK),
        ("BƯỚC 2", "Số hóa & Phân tách câu hỏi với Antigravity + MCP", "Sử dụng Antigravity IDE đọc dữ liệu thô, tách câu đơn và nhóm câu chùm; kết hợp MCP Tavily tra cứu tài liệu và giải bổ sung các câu khuyết đáp án (có giáo viên thẩm định).", COLOR_NAVY_BLUE),
        ("BƯỚC 3", "Lập trình & Phát triển Web App tương tác cao", "Lập trình thuật toán trộn đề ngẫu nhiên 50 câu/60p, thiết kế giao diện Split-Screen tối ưu bài đọc dài, tích hợp Bảng vinh danh Top 100 và Gia sư ảo AI.", RGBColor(0, 102, 204)),
        ("BƯỚC 4", "Đóng gói & Xuất bản đám mây miễn phí 24/7", "Đóng gói toàn bộ mã nguồn tĩnh và phát hành đồng thời lên GitHub Pages và Netlify, đảm bảo hệ thống vận hành liên tục 24/7 với chi phí máy chủ 0 đồng.", RGBColor(180, 40, 20)),
        ("BƯỚC 5", "Tích hợp PWA, Mã QR & Hệ thống thống kê kết quả", "Tạo mã QR truy cập nhanh cho điện thoại, hỗ trợ cài đặt ứng dụng Web (PWA) và tích hợp trang Dashboard phân tích điểm mạnh/yếu của học sinh.", COLOR_ACCENT_RED)
    ]

    card_top = Inches(2.1)
    card_height = Inches(1.3)
    card_gap = Inches(0.18)

    for idx, (step_tag, step_title, step_desc, col) in enumerate(steps):
        current_top = card_top + idx * (card_height + card_gap)
        create_card(slide5, Inches(1.0), current_top, Inches(18.0), card_height, bg_color=COLOR_BG_CARD, border_color=COLOR_BORDER_CARD)
        
        tag_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), current_top + Inches(0.2), Inches(1.8), Inches(0.9))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = col
        tag_box.line.fill.background()
        ttf = tag_box.text_frame
        ttf.margin_top = Inches(0.25)
        tp = ttf.paragraphs[0]
        tp.text = step_tag
        tp.alignment = PP_ALIGN.CENTER
        set_font(tp.runs[0], size_pt=16, bold=True, color=COLOR_TEXT_WHITE)

        c_box = slide5.shapes.add_textbox(Inches(3.2), current_top + Inches(0.12), Inches(15.5), Inches(1.1))
        ctf = c_box.text_frame
        ctf.word_wrap = True
        p1 = ctf.paragraphs[0]
        p1.line_spacing = 1.15
        r1 = p1.add_run()
        r1.text = step_title
        set_font(r1, size_pt=16, bold=True, color=COLOR_NAVY_DARK)
        
        p2 = ctf.add_paragraph()
        p2.line_spacing = 1.2
        r2 = p2.add_run()
        r2.text = step_desc
        set_font(r2, size_pt=14, bold=False, color=COLOR_TEXT_MAIN)

    # ====================================================
    # SLIDE 6 (mới): TRIỂN KHAI THỰC TẾ
    # ====================================================
    print("Formatting Slide 6 (Triển khai thực tế)...")
    slide6 = prs.slides[5]
    for s in slide6.shapes:
        if s.name == 'TextBox 2':
            s.width = Inches(18.0) # ĐẢM BẢO KHÔNG BỊ CỤT TIÊU ĐỀ
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "TRIỂN KHAI THỰC TẾ TẠI TRƯỜNG THPT CHƯƠNG MỸ A"
            set_font(p.runs[0], size_pt=36, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 3':
            clear_shape_text(s)

    # Banner trên
    create_card(slide6, Inches(1.0), Inches(2.1), Inches(18.0), Inches(1.5), bg_color=RGBColor(240, 246, 255), border_color=COLOR_NAVY_BLUE)
    tb_s6_top = slide6.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(17.6), Inches(1.3))
    tf_s6_t = tb_s6_top.text_frame
    tf_s6_t.word_wrap = True
    
    tp1 = tf_s6_t.paragraphs[0]
    tp1.line_spacing = 1.2
    tr1 = tp1.add_run()
    tr1.text = "• Phạm vi thử nghiệm: "
    set_font(tr1, size_pt=15.5, bold=True, color=COLOR_NAVY_DARK)
    tr2 = tp1.add_run()
    tr2.text = "Trực tiếp tại 2 lớp 12A2 và 12A7 - Trường THPT Chương Mỹ A (Hà Nội)."
    set_font(tr2, size_pt=15.5, bold=False, color=COLOR_TEXT_MAIN)
    
    tp2 = tf_s6_t.add_paragraph()
    tp2.line_spacing = 1.2
    tr3 = tp2.add_run()
    tr3.text = "• Thời gian & Quy mô: "
    set_font(tr3, size_pt=15.5, bold=True, color=COLOR_ACCENT_RED)
    tr4 = tp2.add_run()
    tr4.text = "Tháng 10/2026 (Giai đoạn cao điểm ôn thi ĐGNL) | Khảo nghiệm bước đầu với N = 32 học sinh tích cực."
    set_font(tr4, size_pt=15.5, bold=False, color=COLOR_TEXT_MAIN)

    # 3 Khối quy trình
    steps_s6 = [
        ("BƯỚC 1: TRUY CẬP TỰ DO (0Đ CHI PHÍ)", 
         "Quét mã QR trên điện thoại hoặc máy tính cá nhân để vào Web App ngay ngoài giờ học.\n\nKhông cần tài khoản phức tạp, không mất bất kỳ chi phí nào, học sinh có thể ôn luyện mọi lúc mọi nơi.",
         COLOR_NAVY_DARK),
        ("BƯỚC 2: TỰ CHỌN CHẾ ĐỘ ÔN LUYỆN", 
         "Thi thử tổng hợp: 50 câu/60 phút bấm giờ tự động chuẩn format đề thi thật ĐHQGHN.\n\nLuyện chuyên sâu từng dạng: Tiếng Việt, Quy tắc ngữ pháp, Từ vựng, Đọc hiểu văn bản nghệ thuật và thông tin.",
         COLOR_NAVY_BLUE),
        ("BƯỚC 3: TỰ ĐÁNH GIÁ & SỬA LỖI VỚI AI", 
         "Hệ thống tự động chấm điểm tức thì ngay sau khi nộp bài và lưu vào lịch sử cá nhân.\n\nĐọc đáp án chi tiết và tương tác 1-1 với Gia sư AI ở từng câu sai để hiểu bản chất vì sao câu đó sai.",
         COLOR_ACCENT_RED)
    ]

    col_w = Inches(5.7)
    col_gap = Inches(0.45)
    for i, (stitle, sbody, scol) in enumerate(steps_s6):
        c_left = Inches(1.0) + i * (col_w + col_gap)
        create_card(slide6, c_left, Inches(3.9), col_w, Inches(5.3))
        add_header_box(slide6, c_left, Inches(3.9), col_w, Inches(0.65), stitle, bg_color=scol)
        
        tb = slide6.shapes.add_textbox(c_left + Inches(0.2), Inches(4.7), col_w - Inches(0.4), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sbody
        p.line_spacing = 1.3
        set_font(p.runs[0], size_pt=15, bold=False, color=COLOR_TEXT_MAIN)

    # ====================================================
    # SLIDE 7 (mới): KẾT QUẢ VÀ TÁC ĐỘNG
    # ====================================================
    print("Formatting Slide 7 (Kết quả và tác động)...")
    slide7 = prs.slides[6]
    for s in slide7.shapes:
        if s.name == 'TextBox 2':
            s.width = Inches(18.0)
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "KẾT QUẢ VÀ TÁC ĐỘNG THỰC TIỄN"
            set_font(p.runs[0], size_pt=38, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 3':
            clear_shape_text(s)

    # Cột Trái: Tác động thực tế
    create_card(slide7, Inches(1.0), Inches(2.1), Inches(9.2), Inches(7.1))
    add_header_box(slide7, Inches(1.0), Inches(2.1), Inches(9.2), Inches(0.55), "1. TÁC ĐỘNG THỰC TẾ & CHỈ SỐ TIÊU BIỂU", bg_color=COLOR_NAVY_DARK)
    
    tb_s7_l = slide7.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(8.8), Inches(6.2))
    tf_s7_l = tb_s7_l.text_frame
    tf_s7_l.word_wrap = True
    
    impacts = [
        ("Đối với Học sinh | Cải thiện +24,7% điểm số: ", "Có \"Gia sư 24/7\" đồng hành ngoài giờ học; điểm trung bình tăng rõ rệt từ 29,5/50 (Lần 1) lên 36,8/50 (Lần 3+). Ghi nhận 78 lượt làm đề hoàn chỉnh và 145 lượt hỏi Gia sư AI để phân tích câu sai."),
        ("Đối với Giáo viên | Tiết kiệm 98% thời gian: ", "Giải phóng giáo viên khỏi khâu trộn đề thủ công và chấm thi; phát hiện ngay điểm nghẽn kiến thức ở phần Tiếng Việt & Ngữ pháp (tỷ lệ đúng ban đầu chỉ 58%) để kịp thời củng cố trên lớp."),
        ("Bình đẳng cơ hội giáo dục: ", "100% học sinh 12A2 & 12A7 được tiếp cận kho đề chất lượng cao hoàn toàn miễn phí, xóa bỏ hoàn toàn rào cản chi phí luyện thi trung tâm.")
    ]
    for i, (head, body) in enumerate(impacts):
        p = tf_s7_l.paragraphs[0] if i == 0 else tf_s7_l.add_paragraph()
        p.line_spacing = 1.25
        p.space_after = Pt(12)
        r0 = p.add_run()
        r0.text = "• "
        set_font(r0, size_pt=15.5, bold=True, color=COLOR_NAVY_BLUE)
        r1 = p.add_run()
        r1.text = head
        set_font(r1, size_pt=15.5, bold=True, color=COLOR_NAVY_DARK)
        r2 = p.add_run()
        r2.text = body
        set_font(r2, size_pt=14.5, bold=False, color=COLOR_TEXT_MAIN)

    # Cột Phải: Phản hồi thực tế
    create_card(slide7, Inches(10.6), Inches(2.1), Inches(8.4), Inches(7.1))
    add_header_box(slide7, Inches(10.6), Inches(2.1), Inches(8.4), Inches(0.55), "2. PHẢN HỒI THỰC TẾ & MỨC ĐỘ HÀI LÒNG", bg_color=COLOR_ACCENT_RED)
    
    tb_s7_r = slide7.shapes.add_textbox(Inches(10.8), Inches(2.8), Inches(8.0), Inches(6.2))
    tf_s7_r = tb_s7_r.text_frame
    tf_s7_r.word_wrap = True
    
    p_rate = tf_s7_r.paragraphs[0]
    p_rate.line_spacing = 1.2
    p_rate.space_after = Pt(14)
    rr1 = p_rate.add_run()
    rr1.text = "Độ hài lòng chung: 4.9 / 5.0 (Rất hài lòng)\n"
    set_font(rr1, size_pt=17, bold=True, color=COLOR_ACCENT_RED)
    rr2 = p_rate.add_run()
    rr2.text = "(Dựa trên kết quả khảo sát ẩn danh 32 học sinh thử nghiệm)"
    set_font(rr2, size_pt=13.5, italic=True, color=COLOR_TEXT_MUTED)

    quotes = [
        ("Học sinh Lớp 12A2:", "\"Giao diện chia đôi màn hình rất dễ đọc văn bản dài, không bị mỏi mắt như Google Form. Khi làm sai bấm hỏi Gia sư AI được giải thích cặn kẽ từng câu chữ như có cô giáo kèm 1-1 bên cạnh.\""),
        ("Giáo viên Tổ Ngữ văn:", "\"Tận dụng rất tốt kho đề chuẩn của tổ, trộn đề tự động không lo AI bịa kiến thức. Báo cáo thống kê chỉ rõ phần Tiếng Việt học sinh hay sai giúp giáo viên giảng bổ trợ cực kỳ sát thực tế.\"")
    ]
    for head, body in quotes:
        p = tf_s7_r.add_paragraph()
        p.line_spacing = 1.25
        p.space_after = Pt(12)
        r1 = p.add_run()
        r1.text = f"• {head}\n"
        set_font(r1, size_pt=15.5, bold=True, color=COLOR_NAVY_DARK)
        r2 = p.add_run()
        r2.text = body
        set_font(r2, size_pt=14.5, italic=True, color=COLOR_TEXT_MAIN)

    # ====================================================
    # SLIDE 8 (mới): THẢO LUẬN VÀ TỔNG KẾT
    # ====================================================
    print("Formatting Slide 8 (Thảo luận và tổng kết)...")
    slide8 = prs.slides[7]
    for s in slide8.shapes:
        if s.name == 'TextBox 2':
            s.width = Inches(18.0)
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "THẢO LUẬN VÀ TỔNG KẾT"
            set_font(p.runs[0], size_pt=38, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 3':
            clear_shape_text(s)

    conclusions = [
        ("ƯU THẾ VƯỢT TRỘI", 
         "• Khắc phục hoàn toàn nhược điểm cuộn trang mỏi mắt của Google Forms nhờ giao diện chia đôi màn hình (Split-Screen).\n\n• Loại trừ 100% rủi ro bịa kiến thức và gian lận lộ đáp án sớm của ChatGPT tự do.\n\n• Cá nhân hóa tối đa lộ trình sửa sai cho từng học sinh với Gia sư ảo AI.",
         COLOR_NAVY_DARK),
        ("TÍNH BỀN VỮNG & LAN TỎA", 
         "• Nội bộ nhà trường: Lan tỏa từ 2 lớp điểm ra toàn bộ 600+ học sinh khối 12 THPT Chương Mỹ A.\n\n• Cộng đồng: Nhờ nền tảng Web App mở, học sinh lớp 12 trên toàn quốc đều có thể truy cập ôn luyện 0đ.\n\n• Chi phí vận hành = 0 đồng: Lưu trữ vĩnh viễn trên GitHub Pages & Netlify, không phát sinh chi phí server.",
         COLOR_NAVY_BLUE),
        ("ĐỊNH HƯỚNG TƯƠNG LAI", 
         "• Mở rộng ngân hàng đề HSA hàng năm theo sát ma trận đề minh họa mới nhất của ĐHQG Hà Nội.\n\n• Nhân rộng kiến trúc mã nguồn mở sang các môn thi trắc nghiệm khác như Lịch sử, Địa lý, Tiếng Anh.\n\n• Tích hợp tính năng phân tích phổ điểm dự đoán cơ hội đỗ đại học.",
         COLOR_ACCENT_RED)
    ]

    for i, (ctitle, cbody, ccol) in enumerate(conclusions):
        c_left = Inches(1.0) + i * (col_w + col_gap)
        create_card(slide8, c_left, Inches(2.2), col_w, Inches(7.0))
        add_header_box(slide8, c_left, Inches(2.2), col_w, Inches(0.65), ctitle, bg_color=ccol)
        
        tb = slide8.shapes.add_textbox(c_left + Inches(0.2), Inches(3.0), col_w - Inches(0.4), Inches(6.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cbody
        p.line_spacing = 1.3
        set_font(p.runs[0], size_pt=14.5, bold=False, color=COLOR_TEXT_MAIN)

    # ====================================================
    # SLIDE 9 (mới): TÀI LIỆU ĐÍNH KÈM
    # ====================================================
    print("Formatting Slide 9 (Tài liệu đính kèm)...")
    slide9 = prs.slides[8]
    for s in slide9.shapes:
        if s.name == 'TextBox 2':
            s.width = Inches(18.0)
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "TÀI LIỆU VÀ ĐƯỜNG DẪN TRẢI NGHIỆM SẢN PHẨM"
            set_font(p.runs[0], size_pt=36, bold=True, color=COLOR_NAVY_DARK)
        elif s.name == 'TextBox 3':
            clear_shape_text(s)

    # Khung trái
    create_card(slide9, Inches(1.0), Inches(2.2), Inches(11.5), Inches(7.0))
    add_header_box(slide9, Inches(1.0), Inches(2.2), Inches(11.5), Inches(0.55), "HỆ THỐNG ĐƯỜNG DẪN & NGUỒN HỌC LIỆU MỞ", bg_color=COLOR_NAVY_DARK)
    
    tb_s9_l = slide9.shapes.add_textbox(Inches(1.2), Inches(2.9), Inches(11.1), Inches(6.1))
    tf_s9_l = tb_s9_l.text_frame
    tf_s9_l.word_wrap = True
    
    links_info = [
        ("1. Link sản phẩm chính thức (GitHub Pages):", 
         "https://wti200816-bit.github.io/hsa-qualitative-exam-app/", 
         "Hạ tầng máy chủ GitHub bảo mật cao, hoạt động 24/7 ổn định và tốc độ tải trang cực nhanh."),
        ("2. Link dự phòng (Netlify Cloud):", 
         "https://transcendent-cuchufli-30f705.netlify.app", 
         "Hệ thống sao lưu phân tán toàn cầu, tự động đồng bộ liên tục khi cập nhật kho đề mới."),
        ("3. Nguồn học liệu & Bản quyền đề thi:", 
         "Ngân hàng đề chuẩn hóa của GV Tổ Ngữ văn – Ngoại ngữ (THPT Chương Mỹ A)", 
         "Biên soạn chuẩn theo khung ma trận ĐGNL ĐHQGHN và tham khảo nhóm chuyên gia dự án ĐGNL.")
    ]
    
    for i, (head, link_txt, desc) in enumerate(links_info):
        p1 = tf_s9_l.paragraphs[0] if i == 0 else tf_s9_l.add_paragraph()
        p1.line_spacing = 1.2
        r1 = p1.add_run()
        r1.text = f"{head}\n"
        set_font(r1, size_pt=15.5, bold=True, color=COLOR_NAVY_DARK)
        
        r2 = p1.add_run()
        r2.text = f"• {link_txt}\n"
        set_font(r2, size_pt=14.5, bold=True, color=COLOR_ACCENT_RED)
        
        r3 = p1.add_run()
        r3.text = f"   ({desc})\n"
        set_font(r3, size_pt=13.5, italic=True, color=COLOR_TEXT_MUTED)
        p1.space_after = Pt(14)

    # Khung phải: 2 Mã QR Code
    create_card(slide9, Inches(13.0), Inches(2.2), Inches(6.0), Inches(7.0))
    add_header_box(slide9, Inches(13.0), Inches(2.2), Inches(6.0), Inches(0.55), "QUÉT MÃ TRẢI NGHIỆM NGAY", bg_color=COLOR_ACCENT_RED)

    # Chèn QR 1: Netlify Web App
    if os.path.exists(qr_webapp):
        slide9.shapes.add_picture(qr_webapp, Inches(14.8), Inches(2.9), width=Inches(2.4), height=Inches(2.4))
        tb_q1 = slide9.shapes.add_textbox(Inches(13.2), Inches(5.35), Inches(5.6), Inches(0.5))
        qp1 = tb_q1.text_frame.paragraphs[0]
        qp1.text = "Trải nghiệm trực tiếp Web App"
        qp1.alignment = PP_ALIGN.CENTER
        set_font(qp1.runs[0], size_pt=14, bold=True, color=COLOR_NAVY_DARK)

    # Chèn QR 2: GitHub Pages
    if os.path.exists(qr_github):
        slide9.shapes.add_picture(qr_github, Inches(14.8), Inches(5.9), width=Inches(1.8), height=Inches(1.8))
        tb_q2 = slide9.shapes.add_textbox(Inches(13.2), Inches(7.75), Inches(5.6), Inches(0.5))
        qp2 = tb_q2.text_frame.paragraphs[0]
        qp2.text = "Link GitHub Pages & Mã nguồn"
        qp2.alignment = PP_ALIGN.CENTER
        set_font(qp2.runs[0], size_pt=13, bold=True, color=COLOR_TEXT_MUTED)

    # ====================================================
    # LƯU FILE HOÀN THIỆN
    # ====================================================
    print(f"Saving final presentation to: {output_pptx}")
    prs.save(output_pptx)
    print("SUCCESS: Presentation built and saved successfully!")

if __name__ == '__main__':
    src = r'E:\AI\dự thi.pptx'
    clean_src = r'C:\Users\IASNNNN\Documents\antigravity\agitated-galileo\.git\objects\4b\dự thi.pptx'
    temp_work = r'E:\AI\dự thi_working.pptx'
    out_hoanthien = r'E:\AI\dự thi_hoanthien.pptx'
    
    if os.path.exists(clean_src):
        shutil.copy2(clean_src, temp_work)
    else:
        shutil.copy2(src, temp_work)
        
    build_presentation(temp_work, out_hoanthien)
    
    # Ghi đè vào dự thi.pptx để người dùng mở trực tiếp
    shutil.copy2(out_hoanthien, src)
    if os.path.exists(temp_work):
        os.remove(temp_work)
    print("Done! dự thi.pptx has been updated with polished formatting!")
