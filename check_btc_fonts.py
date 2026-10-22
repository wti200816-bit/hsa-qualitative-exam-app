import sys, os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'C:\Users\IASNNNN\Downloads\Mẫu Slide 2026.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"=== Slide {idx+1} ===")
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        print(f"  [{r.font.name}, {r.font.size.pt if r.font.size else 'None'}pt, bold={r.font.bold}]: {r.text[:40]}")
                        break
