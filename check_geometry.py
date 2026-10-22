import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\AI\dự thi.pptx')

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for j, shape in enumerate(slide.shapes):
        pos = f"L={shape.left/914400:.2f} T={shape.top/914400:.2f} W={shape.width/914400:.2f} H={shape.height/914400:.2f}"
        txt = ""
        if shape.has_text_frame:
            txt = shape.text_frame.text[:60].replace("\n", " ")
        print(f"Shape [{j}] ID={shape.shape_id} Name='{shape.name}' Type={shape.shape_type} {pos} | Text='{txt}'")
