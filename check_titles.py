import sys, os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\AI\dự thi.pptx')
print(f"Number of slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    title = ""
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text:
            lines = [l.strip() for l in s.text_frame.text.split("\n") if l.strip()]
            if lines and not title:
                title = lines[0]
    print(f"Slide {i+1}: shapes={len(slide.shapes)}, title='{title}'")
