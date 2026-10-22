import sys, os
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'E:\AI\dự thi.pptx')
slide2 = prs.slides[1]

for idx, s in enumerate(slide2.shapes):
    print(f"Shape {idx}: ID={s.shape_id}, Name='{s.name}', Type={s.shape_type}")
    print(f"   L={s.left/914400:.2f}, T={s.top/914400:.2f}, W={s.width/914400:.2f}, H={s.height/914400:.2f}")
    if s.has_text_frame:
        print(f"   Text: '{s.text_frame.text.strip()}'")
